"""Deterministic tender document inventory, extraction provenance, and readiness.

This module deliberately uses the existing per-bid filesystem persistence.  It
does not decide what a clause means; it records exactly what was (and was not)
read before a later workflow makes that decision.
"""
from __future__ import annotations

import hashlib
import json
import logging
import mimetypes
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from enum import Enum
from pathlib import Path
from typing import Optional

import pdfplumber

logger = logging.getLogger("anamac.tools.tender_documents")

MANIFEST_FILENAME = "document_manifest.json"


class ExtractionStatus(str, Enum):
    DISCOVERED = "DISCOVERED"
    DOWNLOADED = "DOWNLOADED"
    EXTRACTED = "EXTRACTED"
    PARTIALLY_EXTRACTED = "PARTIALLY_EXTRACTED"
    UNSUPPORTED = "UNSUPPORTED"
    FAILED = "FAILED"


class TenderReadiness(str, Enum):
    READY = "READY"
    INCOMPLETE = "INCOMPLETE"
    BLOCKED = "BLOCKED"


@dataclass
class ExtractedPage:
    """Text provenanced to one 1-indexed source PDF page."""

    page_number: int
    text: str


@dataclass
class ExtractedElement:
    """A non-PDF source element with deterministic, source-native provenance."""

    element_type: str
    order: int
    text: str
    block_index: Optional[int] = None
    table_number: Optional[int] = None
    sheet_name: Optional[str] = None
    row_number: Optional[int] = None
    cell_range: Optional[str] = None
    is_heading: bool = False
    is_boq_like: bool = False


@dataclass
class TenderDocument:
    """One discovered tender package document and its extraction outcome."""

    document_id: str
    bid_number: str
    original_filename: str
    stored_path: str
    source_url: Optional[str]
    document_role: str
    mime_type: str
    file_size_bytes: Optional[int] = None
    sha256: Optional[str] = None
    discovered_at: Optional[str] = None
    downloaded_at: Optional[str] = None
    extraction_status: str = ExtractionStatus.DISCOVERED.value
    parser_used: Optional[str] = None
    page_count: Optional[int] = None
    extracted_page_count: int = 0
    element_count: int = 0
    sheet_count: int = 0
    row_count: int = 0
    extracted_at: Optional[str] = None
    extraction_warnings: list[str] = field(default_factory=list)
    extraction_error: Optional[str] = None
    suitable_for_downstream_analysis: bool = False
    pages: list[ExtractedPage] = field(default_factory=list)
    elements: list[ExtractedElement] = field(default_factory=list)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _mime_type(path: Path) -> str:
    known = {
        ".pdf": "application/pdf",
        ".doc": "application/msword",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xls": "application/vnd.ms-excel",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".zip": "application/zip",
        ".txt": "text/plain",
        ".csv": "text/csv",
    }
    return known.get(path.suffix.lower(), mimetypes.guess_type(path.name)[0] or "application/octet-stream")


def _document_id(bid_number: str, stored_path: str) -> str:
    """Stable across repeated processing of the same bid-relative path."""
    return hashlib.sha256(f"{bid_number}:{stored_path}".encode("utf-8")).hexdigest()[:24]


@dataclass
class DocumentManifest:
    bid_number: str
    documents: list[TenderDocument] = field(default_factory=list)
    manifest_version: int = 2
    created_at: str = field(default_factory=_now)
    updated_at: str = field(default_factory=_now)

    @property
    def path(self) -> Path:
        raise AttributeError("A manifest path is supplied to load/save, not serialized")

    def _find(self, stored_path: str) -> Optional[TenderDocument]:
        return next((doc for doc in self.documents if doc.stored_path == stored_path), None)

    def discover(
        self,
        stored_path: str,
        original_filename: str,
        source_url: Optional[str],
        document_role: str,
    ) -> TenderDocument:
        document = self._find(stored_path)
        if document:
            if source_url and not document.source_url:
                document.source_url = source_url
            return document
        document = TenderDocument(
            document_id=_document_id(self.bid_number, stored_path),
            bid_number=self.bid_number,
            original_filename=original_filename,
            stored_path=stored_path,
            source_url=source_url,
            document_role=document_role,
            mime_type=_mime_type(Path(original_filename)),
            discovered_at=_now(),
        )
        self.documents.append(document)
        return document

    def record_download(self, document: TenderDocument, bid_dir: Path, success: bool, error: Optional[str] = None) -> None:
        if not success:
            document.extraction_status = ExtractionStatus.FAILED.value
            document.extraction_error = error or "Download failed"
            document.suitable_for_downstream_analysis = False
            return
        file_path = bid_dir / document.stored_path
        if not file_path.is_file():
            self.record_download(document, bid_dir, False, "Downloaded file is missing")
            return
        document.file_size_bytes = file_path.stat().st_size
        document.sha256 = _file_hash(file_path)
        document.downloaded_at = _now()
        document.mime_type = _mime_type(file_path)
        document.extraction_status = ExtractionStatus.DOWNLOADED.value
        document.extraction_error = None
        document.extraction_warnings = []
        document.suitable_for_downstream_analysis = False

    def extract(self, document: TenderDocument, bid_dir: Path) -> None:
        """Dispatch safe, deterministic extraction without inventing source provenance."""
        file_path = bid_dir / document.stored_path
        document.pages = []
        document.elements = []
        document.extracted_page_count = 0
        document.element_count = 0
        document.sheet_count = 0
        document.row_count = 0
        document.extraction_warnings = []
        document.extraction_error = None
        document.extracted_at = _now()

        if not file_path.is_file():
            document.extraction_status = ExtractionStatus.FAILED.value
            document.extraction_error = "Stored file is missing"
            document.suitable_for_downstream_analysis = False
            return

        suffix = file_path.suffix.lower()
        parser = _PARSERS.get(suffix)
        if not parser:
            document.extraction_status = ExtractionStatus.UNSUPPORTED.value
            document.parser_used = None
            document.extraction_error = "No deterministic parser is configured for this file type"
            document.suitable_for_downstream_analysis = False
            return
        parser(document, file_path)

    def counts(self) -> dict[str, int]:
        statuses = [doc.extraction_status for doc in self.documents]
        return {
            "discovered": len(self.documents),
            "downloaded": sum(doc.downloaded_at is not None for doc in self.documents),
            "extracted": statuses.count(ExtractionStatus.EXTRACTED.value),
            "partially_extracted": statuses.count(ExtractionStatus.PARTIALLY_EXTRACTED.value),
            "unsupported": statuses.count(ExtractionStatus.UNSUPPORTED.value),
            "failed": statuses.count(ExtractionStatus.FAILED.value),
            "expected_pages": sum(doc.page_count or 0 for doc in self.documents),
            "extracted_pages": sum(doc.extracted_page_count for doc in self.documents),
        }

    def readiness(self) -> dict[str, object]:
        primary = next((doc for doc in self.documents if doc.document_role == "PRIMARY"), None)
        if not primary:
            return {"status": TenderReadiness.BLOCKED.value, "reasons": ["Primary tender document was not discovered"], "affected_documents": []}
        if primary.extraction_status in {
            ExtractionStatus.DISCOVERED.value,
            ExtractionStatus.DOWNLOADED.value,
            ExtractionStatus.UNSUPPORTED.value,
            ExtractionStatus.FAILED.value,
        }:
            return {
                "status": TenderReadiness.BLOCKED.value,
                "reasons": [f"Primary tender document is {primary.extraction_status}"],
                "affected_documents": [_readiness_document(primary)],
            }

        incomplete = [
            doc for doc in self.documents
            if doc.extraction_status != ExtractionStatus.EXTRACTED.value
        ]
        if incomplete:
            return {
                "status": TenderReadiness.INCOMPLETE.value,
                "reasons": [f"{doc.stored_path}: {doc.extraction_status}" for doc in incomplete],
                "affected_documents": [_readiness_document(doc) for doc in incomplete],
            }
        return {"status": TenderReadiness.READY.value, "reasons": [], "affected_documents": []}

    def to_dict(self) -> dict:
        return {
            "manifest_version": self.manifest_version,
            "bid_number": self.bid_number,
            "created_at": self.created_at,
            "updated_at": self.updated_at,
            "counts": self.counts(),
            "readiness": self.readiness(),
            "documents": [asdict(document) for document in self.documents],
        }


_MAX_OFFICE_FILE_BYTES = 50 * 1024 * 1024
_MAX_ARCHIVE_MEMBERS = 2_000
_MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024


def _readiness_document(document: TenderDocument) -> dict[str, object]:
    return {
        "document_id": document.document_id,
        "stored_path": document.stored_path,
        "original_filename": document.original_filename,
        "document_role": document.document_role,
        "extraction_status": document.extraction_status,
        "warnings": document.extraction_warnings,
        "error": document.extraction_error,
        "corrective_action": "Replace or re-download the document and run extraction again.",
    }


def _check_office_archive(path: Path) -> Optional[str]:
    """Bound ZIP-backed Office parsing before a library expands the archive."""
    if path.stat().st_size > _MAX_OFFICE_FILE_BYTES:
        return f"Office document exceeds {_MAX_OFFICE_FILE_BYTES // (1024 * 1024)} MB extraction limit"
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > _MAX_ARCHIVE_MEMBERS:
                return "Office archive has too many members"
            if sum(info.file_size for info in infos) > _MAX_UNCOMPRESSED_BYTES:
                return "Office archive exceeds uncompressed extraction limit"
    except (zipfile.BadZipFile, OSError) as exc:
        return f"Office archive is unreadable: {exc}"
    return None


def _finish_elements(document: TenderDocument, empty_error: str) -> None:
    document.element_count = len(document.elements)
    if document.element_count:
        document.extraction_status = (
            ExtractionStatus.PARTIALLY_EXTRACTED.value
            if document.extraction_warnings
            else ExtractionStatus.EXTRACTED.value
        )
        document.suitable_for_downstream_analysis = not document.extraction_warnings
        return
    document.extraction_status = ExtractionStatus.FAILED.value
    document.extraction_error = empty_error
    document.suitable_for_downstream_analysis = False


def _extract_pdf(document: TenderDocument, path: Path) -> None:
    document.parser_used = "pdfplumber"
    try:
        with pdfplumber.open(str(path)) as pdf:
            document.page_count = len(pdf.pages)
            for page_number, page in enumerate(pdf.pages, start=1):
                try:
                    page_text = (page.extract_text() or "").strip()
                except Exception as exc:
                    document.extraction_warnings.append(f"Page {page_number}: {exc}")
                    continue
                if not page_text:
                    document.extraction_warnings.append(f"Page {page_number}: no extractable text")
                    continue
                document.pages.append(ExtractedPage(page_number=page_number, text=page_text))
    except Exception as exc:
        document.extraction_status = ExtractionStatus.FAILED.value
        document.extraction_error = f"PDF extraction failed: {exc}"
        return
    document.extracted_page_count = len(document.pages)
    if document.page_count and document.extracted_page_count == document.page_count:
        document.extraction_status = ExtractionStatus.EXTRACTED.value
        document.suitable_for_downstream_analysis = True
    elif document.extracted_page_count:
        document.extraction_status = ExtractionStatus.PARTIALLY_EXTRACTED.value
    else:
        document.extraction_status = ExtractionStatus.FAILED.value
        document.extraction_error = "PDF contains no extractable text"


def _extract_docx(document: TenderDocument, path: Path) -> None:
    document.parser_used = "python-docx"
    if error := _check_office_archive(path):
        document.extraction_status, document.extraction_error = ExtractionStatus.FAILED.value, error
        return
    try:
        from docx import Document
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        docx_document = Document(str(path))
        order = 0
        table_number = 0
        for block_index, child in enumerate(docx_document.element.body.iterchildren(), start=1):
            if isinstance(child, CT_P):
                paragraph = Paragraph(child, docx_document)
                text = paragraph.text.strip()
                if text:
                    order += 1
                    style_name = paragraph.style.name if paragraph.style else ""
                    document.elements.append(ExtractedElement(
                        element_type="heading" if style_name.lower().startswith("heading") else "paragraph",
                        order=order, block_index=block_index, text=text,
                        is_heading=style_name.lower().startswith("heading"),
                    ))
            elif isinstance(child, CT_Tbl):
                table_number += 1
                table = Table(child, docx_document)
                rows = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        rows.append(" | ".join(cells))
                if rows:
                    order += 1
                    document.elements.append(ExtractedElement(
                        element_type="table", order=order, block_index=block_index,
                        table_number=table_number, text="\n".join(rows),
                    ))
    except Exception as exc:
        document.extraction_status = ExtractionStatus.FAILED.value
        document.extraction_error = f"DOCX extraction failed: {exc}"
        return
    _finish_elements(document, "DOCX contains no extractable text or table content")


def _is_boq_header(values: list[str]) -> bool:
    joined = " ".join(value.lower() for value in values)
    matches = sum(term in joined for term in ("item", "description", "specification", "quantity", "unit", "rate", "amount"))
    return matches >= 2


def _extract_xlsx(document: TenderDocument, path: Path) -> None:
    document.parser_used = "openpyxl"
    if error := _check_office_archive(path):
        document.extraction_status, document.extraction_error = ExtractionStatus.FAILED.value, error
        return
    try:
        from openpyxl import load_workbook
        # data_only=False deliberately retains formulas as literal data and never evaluates them.
        workbook = load_workbook(filename=str(path), read_only=True, data_only=False, keep_links=False)
        document.sheet_count = len(workbook.worksheets)
        order = 0
        for worksheet in workbook.worksheets:
            nonempty_rows: list[tuple[int, list[str], str]] = []
            for row_number, row in enumerate(worksheet.iter_rows(), start=1):
                cells = []
                coordinates = []
                for cell in row:
                    if cell.value is not None:
                        cells.append(str(cell.value))
                        coordinates.append(cell.coordinate)
                if cells:
                    cell_range = f"{coordinates[0]}:{coordinates[-1]}" if len(coordinates) > 1 else coordinates[0]
                    nonempty_rows.append((row_number, cells, cell_range))
            document.row_count += len(nonempty_rows)
            boq_like = bool(nonempty_rows and _is_boq_header(nonempty_rows[0][1]))
            for row_number, cells, cell_range in nonempty_rows:
                order += 1
                document.elements.append(ExtractedElement(
                    element_type="worksheet_row", order=order, text=" | ".join(cells),
                    sheet_name=worksheet.title, row_number=row_number, cell_range=cell_range,
                    is_boq_like=boq_like,
                ))
        workbook.close()
    except Exception as exc:
        document.extraction_status = ExtractionStatus.FAILED.value
        document.extraction_error = f"XLSX extraction failed: {exc}"
        return
    _finish_elements(document, "XLSX contains no extractable workbook content")


_PARSERS = {".pdf": _extract_pdf, ".docx": _extract_docx, ".xlsx": _extract_xlsx}


def save_manifest(manifest: DocumentManifest, bid_dir: Path) -> Path:
    manifest.updated_at = _now()
    destination = bid_dir / MANIFEST_FILENAME
    temporary = destination.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(manifest.to_dict(), indent=2, ensure_ascii=False), encoding="utf-8")
    temporary.replace(destination)
    return destination


def load_manifest(bid_dir: Path) -> Optional[DocumentManifest]:
    source = bid_dir / MANIFEST_FILENAME
    if not source.exists():
        return None
    data = json.loads(source.read_text(encoding="utf-8"))
    documents = []
    for raw in data.get("documents", []):
        raw["pages"] = [ExtractedPage(**page) for page in raw.get("pages", [])]
        raw["elements"] = [ExtractedElement(**element) for element in raw.get("elements", [])]
        documents.append(TenderDocument(**raw))
    return DocumentManifest(
        bid_number=data["bid_number"],
        documents=documents,
        manifest_version=data.get("manifest_version", 1),
        created_at=data.get("created_at", _now()),
        updated_at=data.get("updated_at", _now()),
    )


def build_manifest_for_existing_bid(bid_number: str, bid_dir: Path) -> DocumentManifest:
    """Create provenance for legacy bid folders without changing their existing consumers."""
    manifest = DocumentManifest(bid_number=bid_number)
    metadata = {}
    metadata_path = bid_dir / "bid_metadata.json"
    if metadata_path.exists():
        try:
            metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        except Exception:
            logger.warning("Could not read bid metadata for legacy manifest: %s", bid_dir)

    primary = bid_dir / "bid_document.pdf"
    if primary.exists():
        doc = manifest.discover("bid_document.pdf", primary.name, metadata.get("detail_url"), "PRIMARY")
        manifest.record_download(doc, bid_dir, True)
        manifest.extract(doc, bid_dir)

    supporting_dir = bid_dir / "supporting_docs"
    if supporting_dir.exists():
        for file_path in sorted(path for path in supporting_dir.iterdir() if path.is_file()):
            stored_path = str(file_path.relative_to(bid_dir))
            doc = manifest.discover(stored_path, file_path.name, None, "SUPPORTING")
            manifest.record_download(doc, bid_dir, True)
            manifest.extract(doc, bid_dir)

    save_manifest(manifest, bid_dir)
    return manifest


def load_or_build_manifest(bid_number: str, bid_dir: Path) -> DocumentManifest:
    return load_manifest(bid_dir) or build_manifest_for_existing_bid(bid_number, bid_dir)


def analysis_readiness(bid_number: str, bid_dir: Path) -> tuple[DocumentManifest, dict[str, object]]:
    """The sole readiness entry point for tender-backed AI workflows."""
    manifest = load_or_build_manifest(bid_number, bid_dir)
    return manifest, manifest.readiness()


def analysis_gate(bid_number: str, bid_dir: Path) -> dict[str, object]:
    """Stable API/agent representation; no route reimplements readiness rules."""
    manifest, readiness = analysis_readiness(bid_number, bid_dir)
    return {
        "bid_number": bid_number,
        "readiness": readiness,
        "blocked": readiness["status"] == TenderReadiness.BLOCKED.value,
        "analysis_incomplete": readiness["status"] == TenderReadiness.INCOMPLETE.value,
        "counts": manifest.counts(),
    }


def analysis_document_text(manifest: DocumentManifest) -> tuple[str, dict[str, str]]:
    """Return only manifest-provenanced tender content, preserving source identity."""
    documents: dict[str, str] = {}
    for document in manifest.documents:
        parts: list[str] = []
        for page in document.pages:
            parts.append(f"[page {page.page_number}]\n{page.text}")
        for element in document.elements:
            provenance = [f"element {element.order}"]
            if element.block_index is not None:
                provenance.append(f"block {element.block_index}")
            if element.table_number is not None:
                provenance.append(f"table {element.table_number}")
            if element.sheet_name:
                provenance.append(f"sheet {element.sheet_name}")
            if element.row_number is not None:
                provenance.append(f"row {element.row_number}")
            if element.cell_range:
                provenance.append(element.cell_range)
            parts.append(f"[{' | '.join(provenance)}]\n{element.text}")
        if parts:
            documents[document.stored_path] = "\n\n".join(parts)
    primary = next((doc for doc in manifest.documents if doc.document_role == "PRIMARY"), None)
    return (documents.get(primary.stored_path, "") if primary else ""), documents
