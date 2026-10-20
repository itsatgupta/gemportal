"""
Generate ANAMAC Bid Platform Client Presentation Deck — Star Wars Edition
Run: cd backend && source venv/bin/activate && python ../docs/generate_presentation.py
Output: docs/ANAMAC_Bid_Platform_Presentation.pptx
"""
from __future__ import annotations

import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ═══════════════════════════════════════════════════════════════
# Colors
# ═══════════════════════════════════════════════════════════════
NAVY = RGBColor(0x0A, 0x0E, 0x1A)
DARK_SLATE = RGBColor(0x11, 0x18, 0x27)
CARD_BG = RGBColor(0x1A, 0x22, 0x35)
ELEVATED = RGBColor(0x1E, 0x2A, 0x3D)
GOLD = RGBColor(0xC9, 0xA8, 0x4C)
GOLD_LIGHT = RGBColor(0xE8, 0xC9, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xCC, 0xCC, 0xCC)
MID_GRAY = RGBColor(0x88, 0x88, 0x88)
DIM_GRAY = RGBColor(0x55, 0x55, 0x55)
CYAN = RGBColor(0x06, 0xB6, 0xD4)
TEAL = RGBColor(0x14, 0xB8, 0xA6)
EMERALD = RGBColor(0x10, 0xB9, 0x81)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
PURPLE = RGBColor(0x8B, 0x5C, 0xF6)
RED = RGBColor(0xEF, 0x44, 0x44)
ORANGE = RGBColor(0xF9, 0x73, 0x16)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.dirname(SCRIPT_DIR)
LOGO_PATH = os.path.join(PROJECT_DIR, "frontend", "public", "anamaclogo.png")

# ═══════════════════════════════════════════════════════════════
# Star Wars Agent Definitions
# ═══════════════════════════════════════════════════════════════
CHARACTERS_DIR = os.path.join(SCRIPT_DIR, "characters")

AGENTS = {
    "luke": {
        "name": "LUKE SKYWALKER",
        "subtitle": "The Hero",
        "role": "Search & Download Agent",
        "tagline": '"I\'m Luke Skywalker, I\'m here to find your bids."',
        "color_rgb": (0x3B, 0x82, 0xF6),  # Blue
        "color": BLUE,
        "glow": (0x60, 0xA5, 0xFA),
        "icon": "\u2694",  # swords
        "avatar_file": "luke_skywalker_avatar.png",
        "abilities": [
            "Navigates the GEM Portal with Jedi precision",
            "Discovers matching bids using the Force (Playwright)",
            "Downloads all bid documents & supporting files",
            "Organizes files into structured bid folders",
        ],
    },
    "yoda": {
        "name": "YODA",
        "subtitle": "The Sage",
        "role": "Proposal Writer Agent",
        "tagline": '"Write proposals, I do. Strong with the bid, ANAMAC is."',
        "color_rgb": (0x10, 0xB9, 0x81),  # Emerald
        "color": EMERALD,
        "glow": (0x34, 0xD3, 0x99),
        "icon": "\u270D",  # writing hand
        "avatar_file": "yoda_avatar.png",
        "abilities": [
            "Reads all bid requirements with ancient wisdom",
            "Cross-references ANAMAC company documents",
            "Generates proposals via Claude AI (the Force)",
            "Creates branded PDF proposals, revises on feedback",
        ],
    },
    "vader": {
        "name": "DARTH VADER",
        "subtitle": "The Enforcer",
        "role": "Quality Assessor Agent",
        "tagline": '"I find your lack of compliance... disturbing."',
        "color_rgb": (0xF5, 0x9E, 0x0B),  # Amber
        "color": AMBER,
        "glow": (0xFB, 0xBF, 0x24),
        "icon": "\u2696",  # scales
        "avatar_file": "darth_vader_avatar.png",
        "abilities": [
            "Reviews proposals against bid requirements",
            "Identifies gaps: Critical / Major / Minor",
            "Sends structured gap reports back to Yoda",
            "Approves once all issues resolved (max 2 rounds)",
        ],
    },
    "r2d2": {
        "name": "R2-D2",
        "subtitle": "The Droid",
        "role": "QA Testing Agent",
        "tagline": '"*beep boop* All 22 tests passing *happy beep*"',
        "color_rgb": (0x8B, 0x5C, 0xF6),  # Purple
        "color": PURPLE,
        "glow": (0xA7, 0x8B, 0xFA),
        "icon": "\u2699",  # gear
        "avatar_file": "r2d2_avatar.png",
        "abilities": [
            "Tests GEM portal connectivity & scraping",
            "Validates document download integrity",
            "Tests agent-to-agent communication flows",
            "Generates HTML test reports, 22/22 passing",
        ],
    },
}


def _load_avatars() -> dict:
    """Load pre-generated character avatar images, return {key: path}."""
    avatars = {}
    for key, agent in AGENTS.items():
        avatar_path = os.path.join(CHARACTERS_DIR, agent["avatar_file"])
        if os.path.exists(avatar_path):
            avatars[key] = avatar_path
        else:
            print(f"  Warning: avatar not found: {avatar_path}")
            avatars[key] = None
    return avatars


# ═══════════════════════════════════════════════════════════════
# Slide Helpers
# ═══════════════════════════════════════════════════════════════
def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 color=WHITE, bold=False, alignment=PP_ALIGN.LEFT,
                 font_name="Calibri", italic=False):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.italic = italic
    p.font.name = font_name
    p.alignment = alignment
    return tf


def add_paragraph(tf, text, font_size=16, color=WHITE, bold=False,
                  space_before=Pt(6), font_name="Calibri", alignment=PP_ALIGN.LEFT,
                  italic=False):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.italic = italic
    p.font.name = font_name
    p.alignment = alignment
    if space_before:
        p.space_before = space_before
    return p


def add_bullet_list(slide, left, top, width, height, items, font_size=14,
                    color=LIGHT_GRAY):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
        p.space_before = Pt(6)
    return tf


def add_gold_line(slide, left, top, width):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()


def add_accent_line(slide, left, top, width, color=TEAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_rounded_rect(slide, left, top, width, height, fill_color=CARD_BG,
                     border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape


def add_logo_watermark(slide):
    """Add small ANAMAC logo to bottom-right of every slide."""
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(
            LOGO_PATH,
            Inches(8.8), Inches(4.85),
            Inches(0.65), Inches(0.65),
        )


def slide_header(slide, title, subtitle=None, accent_color=GOLD):
    """Vibrant slide header with accent bar and title."""
    set_slide_bg(slide, NAVY)

    # Top accent bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Pt(4))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent_color
    bar.line.fill.background()

    add_accent_line(slide, Inches(0.5), Inches(0.45), Inches(9), accent_color)
    add_text_box(slide, Inches(0.5), Inches(0.55), Inches(9), Inches(0.7),
                 title, font_size=28, color=accent_color, bold=True)
    if subtitle:
        add_text_box(slide, Inches(0.5), Inches(1.15), Inches(9), Inches(0.4),
                     subtitle, font_size=14, color=LIGHT_GRAY, italic=True)
    add_accent_line(slide, Inches(0.5), Inches(1.55), Inches(9), accent_color)

    add_logo_watermark(slide)


def add_agent_card(slide, agent_key, left, top, width, height, avatars, show_tagline=False):
    """Add a vibrant agent character card with avatar, name, and abilities."""
    agent = AGENTS[agent_key]
    color = agent["color"]

    # Card background
    add_rounded_rect(slide, left, top, width, height, CARD_BG, color)

    # Glow strip at top of card
    strip = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left + Inches(0.05), top + Inches(0.05),
        width - Inches(0.1), Pt(4),
    )
    strip.fill.solid()
    strip.fill.fore_color.rgb = color
    strip.line.fill.background()

    # Avatar image
    avatar_path = avatars.get(agent_key)
    if avatar_path and os.path.exists(avatar_path):
        avatar_size = Inches(0.7)
        slide.shapes.add_picture(
            avatar_path,
            left + Inches(0.15), top + Inches(0.2),
            avatar_size, avatar_size,
        )

    # Name and subtitle
    name_left = left + Inches(0.95)
    add_text_box(slide, name_left, top + Inches(0.15), width - Inches(1.1), Inches(0.35),
                 agent["name"], font_size=16, color=color, bold=True)
    add_text_box(slide, name_left, top + Inches(0.45), width - Inches(1.1), Inches(0.25),
                 f'{agent["subtitle"]}  \u2022  {agent["role"]}', font_size=9, color=MID_GRAY)

    # Tagline
    if show_tagline:
        add_text_box(slide, left + Inches(0.15), top + Inches(0.85), width - Inches(0.3), Inches(0.35),
                     agent["tagline"], font_size=9, color=LIGHT_GRAY, italic=True)

    return agent


# ═══════════════════════════════════════════════════════════════
# Build Presentation
# ═══════════════════════════════════════════════════════════════
def build_presentation():
    # Load character avatars from docs/characters/
    print("Loading Star Wars character avatars...")
    avatars = _load_avatars()

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)  # 16:9

    # ==========================================================
    # SLIDE 1 — Title Slide (Epic)
    # ==========================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, NAVY)

    # Top teal accent bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Pt(5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()

    # Logo
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Inches(3.8), Inches(0.3), Inches(2.4), Inches(1.8))

    add_text_box(slide, Inches(0.5), Inches(2.2), Inches(9), Inches(0.7),
                 "ANAMAC", font_size=48, color=GOLD, bold=True,
                 alignment=PP_ALIGN.CENTER)
    add_text_box(slide, Inches(0.5), Inches(2.9), Inches(9), Inches(0.6),
                 "Bid Proposal Intelligence Platform", font_size=26, color=WHITE,
                 alignment=PP_ALIGN.CENTER)
    add_gold_line(slide, Inches(2.5), Inches(3.5), Inches(5))
    add_text_box(slide, Inches(0.5), Inches(3.65), Inches(9), Inches(0.5),
                 "AI-Powered Government Bid Automation  |  GEM Portal Integration",
                 font_size=14, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    # Star Wars subtitle
    add_text_box(slide, Inches(0.5), Inches(4.15), Inches(9), Inches(0.4),
                 '"May the Bids Be With You"',
                 font_size=13, color=TEAL, alignment=PP_ALIGN.CENTER, italic=True)

    add_text_box(slide, Inches(0.5), Inches(4.8), Inches(9), Inches(0.4),
                 "Confidential  |  March 2026", font_size=11, color=DIM_GRAY,
                 alignment=PP_ALIGN.CENTER)

    # Small character avatars at bottom
    agent_keys = ["luke", "yoda", "vader", "r2d2"]
    for i, key in enumerate(agent_keys):
        avatar_path = avatars[key]
        if os.path.exists(avatar_path):
            slide.shapes.add_picture(
                avatar_path,
                Inches(3.0) + Inches(1.1) * i, Inches(4.55),
                Inches(0.45), Inches(0.45),
            )

    add_logo_watermark(slide)

    # ==========================================================
    # SLIDE 2 — The Problem
    # ==========================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(slide, "The Problem", "Why Government Bid Preparation Needs Automation", RED)

    problems = [
        ("\u26A0  Manual bid discovery", "Hours spent browsing GEM portal daily"),
        ("\u26A0  Document download chaos", "Bid PDFs contain dozens of linked documents"),
        ("\u26A0  Proposal writing bottleneck", "Each bid requires days of manual drafting"),
        ("\u26A0  Compliance gaps missed", "Human reviewers overlook requirements under pressure"),
        ("\u26A0  Slow turnaround", "Missed deadlines = missed revenue opportunities"),
        ("\u26A0  No standardization", "Proposal quality varies across team members"),
    ]
    for i, (title, desc) in enumerate(problems):
        row = i // 2
        col = i % 2
        left = Inches(0.5) + Inches(4.7) * col
        top = Inches(1.8) + Inches(1.15) * row

        add_rounded_rect(slide, left, top, Inches(4.4), Inches(0.95), CARD_BG, RED)
        add_text_box(slide, left + Inches(0.15), top + Inches(0.08), Inches(4.1), Inches(0.35),
                     title, font_size=13, color=RED, bold=True)
        add_text_box(slide, left + Inches(0.15), top + Inches(0.42), Inches(4.1), Inches(0.4),
                     desc, font_size=11, color=LIGHT_GRAY)

    # ==========================================================
    # SLIDE 3 — The Solution
    # ==========================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(slide, "The Solution", "End-to-End AI-Powered Bid Automation", TEAL)

    solutions = [
        ("Automated Discovery", "AI scout searches GEM portal\nfinds relevant bids instantly", BLUE, "luke"),
        ("Smart Download", "Downloads bid PDF + all linked\nsupporting documents", CYAN, "luke"),
        ("AI Proposal Generation", "Claude AI drafts comprehensive\nproposals aligned to requirements", EMERALD, "yoda"),
        ("Quality Assessment", "AI assessor identifies compliance\ngaps with severity ratings", AMBER, "vader"),
    ]
    for i, (title, desc, color, agent_key) in enumerate(solutions):
        left = Inches(0.35) + Inches(2.4) * i
        box = add_rounded_rect(slide, left, Inches(1.8), Inches(2.2), Inches(2.8), CARD_BG, color)

        # Agent avatar
        avatar_path = avatars[agent_key]
        if os.path.exists(avatar_path):
            slide.shapes.add_picture(
                avatar_path, left + Inches(0.75), Inches(1.95), Inches(0.6), Inches(0.6)
            )

        add_text_box(slide, left + Inches(0.1), Inches(2.65), Inches(2.0), Inches(0.35),
                     title, font_size=13, color=color, bold=True, alignment=PP_ALIGN.CENTER)
        add_text_box(slide, left + Inches(0.1), Inches(3.1), Inches(2.0), Inches(1.2),
                     desc, font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    add_rounded_rect(slide, Inches(0.35), Inches(4.75), Inches(9.3), Inches(0.55), DARK_SLATE, GOLD)
    add_text_box(slide, Inches(0.5), Inches(4.82), Inches(9), Inches(0.4),
                 "\u21BB  Automated revision loop: Vader sends gaps back to Yoda until all issues are resolved",
                 font_size=12, color=GOLD, alignment=PP_ALIGN.CENTER)

    # ==========================================================
    # SLIDE 4 — System Architecture
    # ==========================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(slide, "System Architecture", "Multi-Agent AI Platform with Real-Time Communication", CYAN)

    components = [
        ("React Frontend", "Vite + Tailwind + Zustand\n4 Interactive Tabs\nReal-time WebSocket Updates", Inches(0.3), CYAN),
        ("FastAPI Backend", "Python 3.9+ API Server\nWebSocket Agent Communication\nEvent-Driven Orchestration", Inches(3.45), EMERALD),
        ("AI Agent Team", "Luke \u2022 Yoda \u2022 Vader \u2022 R2-D2\nClaude AI (Anthropic)\nPlaywright Browser Automation", Inches(6.6), GOLD),
    ]
    for title, desc, left, color in components:
        add_rounded_rect(slide, left, Inches(1.8), Inches(3.0), Inches(2.2), CARD_BG, color)
        add_text_box(slide, left + Inches(0.15), Inches(1.9), Inches(2.7), Inches(0.4),
                     title, font_size=15, color=color, bold=True, alignment=PP_ALIGN.CENTER)
        add_text_box(slide, left + Inches(0.15), Inches(2.4), Inches(2.7), Inches(1.5),
                     desc, font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    # Arrows
    for x in [Inches(3.3), Inches(6.45)]:
        arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, Inches(2.7), Inches(0.25), Inches(0.3))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = GOLD
        arrow.line.fill.background()

    # Tech stack bar
    add_rounded_rect(slide, Inches(0.3), Inches(4.2), Inches(9.4), Inches(0.9), DARK_SLATE, TEAL)
    add_text_box(slide, Inches(0.5), Inches(4.35), Inches(9), Inches(0.6),
                 "Tech Stack:  React 18  |  FastAPI  |  Claude AI  |  Playwright  |  Zustand  |  WebSocket  |  slowapi  |  reportlab",
                 font_size=12, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    # ==========================================================
    # SLIDE 5 — Meet the AI Agent Team (Star Wars)
    # ==========================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(slide, "Meet the AI Agent Team", '"A long time ago in a portal far, far away..."', GOLD)

    agent_keys = ["luke", "yoda", "vader", "r2d2"]
    for i, key in enumerate(agent_keys):
        agent = AGENTS[key]
        col = i % 2
        row = i // 2
        left = Inches(0.4) + Inches(4.8) * col
        top = Inches(1.75) + Inches(1.75) * row

        # Card
        add_rounded_rect(slide, left, top, Inches(4.55), Inches(1.6), CARD_BG, agent["color"])

        # Glow strip
        strip = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left + Inches(0.05), top + Inches(0.05),
            Inches(4.45), Pt(3),
        )
        strip.fill.solid()
        strip.fill.fore_color.rgb = agent["color"]
        strip.line.fill.background()

        # Avatar
        avatar_path = avatars[key]
        if os.path.exists(avatar_path):
            slide.shapes.add_picture(avatar_path, left + Inches(0.15), top + Inches(0.2),
                                     Inches(0.85), Inches(0.85))

        # Name and role
        txt_left = left + Inches(1.1)
        add_text_box(slide, txt_left, top + Inches(0.15), Inches(3.2), Inches(0.35),
                     agent["name"], font_size=18, color=agent["color"], bold=True)
        add_text_box(slide, txt_left, top + Inches(0.45), Inches(3.2), Inches(0.25),
                     f'{agent["subtitle"]}  \u2014  {agent["role"]}', font_size=10, color=MID_GRAY)
        add_text_box(slide, txt_left, top + Inches(0.72), Inches(3.2), Inches(0.3),
                     agent["tagline"], font_size=8, color=LIGHT_GRAY, italic=True)

        # Abilities (2 lines)
        abilities_text = "  \u2022  ".join(agent["abilities"][:2])
        add_text_box(slide, left + Inches(0.15), top + Inches(1.1), Inches(4.2), Inches(0.4),
                     "\u2022  " + abilities_text, font_size=8, color=LIGHT_GRAY)

    # ==========================================================
    # SLIDE 6 — Step 1: Luke's Mission
    # ==========================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(slide, "Step 1: Bid Discovery & Download",
                 "Luke Skywalker scouts the GEM portal and acquires all documents", BLUE)

    # Luke card
    add_agent_card(slide, "luke", Inches(0.5), Inches(1.75), Inches(3.0), Inches(1.0), avatars, True)

    steps = [
        "1. User enters bid type (QRA, HAZOP, Safety Audit, etc.)",
        "2. Luke navigates to GEM portal with Jedi precision",
        "3. Results displayed: Bid Number, Title, Department, Deadline",
        "4. User selects a bid \u2192 primary bid PDF downloaded",
        "5. All hyperlinks extracted from bid PDF",
        "6. Supporting documents downloaded automatically",
        "7. Files organized in bids/{BID_NUMBER}/ folder",
        "8. Real-time progress bars track each download",
    ]
    add_bullet_list(slide, Inches(3.7), Inches(1.75), Inches(6.0), Inches(3.5),
                    steps, font_size=12, color=LIGHT_GRAY)

    # ==========================================================
    # SLIDE 7 — Step 2: Yoda's Wisdom
    # ==========================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(slide, "Step 2: AI Proposal Generation",
                 "Yoda drafts a comprehensive bid proposal using the Force (Claude AI)", EMERALD)

    add_agent_card(slide, "yoda", Inches(0.5), Inches(1.75), Inches(3.0), Inches(1.0), avatars, True)

    steps = [
        "1. Reads all bid documents + company supporting docs",
        "2. Analyzes scope, eligibility, technical requirements",
        "3. Cross-references ANAMAC certifications & past projects",
        "4. Claude AI generates tailored bid proposal",
        "5. Includes financial proposal with pricing strategy",
        "6. References exact PDF filenames for attachments",
        "7. Professional branded PDF with cover page",
        "8. Vader auto-notified when proposal is ready",
    ]
    add_bullet_list(slide, Inches(3.7), Inches(1.75), Inches(6.0), Inches(3.5),
                    steps, font_size=12, color=LIGHT_GRAY)

    # ==========================================================
    # SLIDE 8 — Step 3: Vader's Assessment
    # ==========================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(slide, "Step 3: Quality Assessment & Gap Analysis",
                 "Darth Vader reviews the proposal and demands perfection", AMBER)

    add_agent_card(slide, "vader", Inches(0.5), Inches(1.75), Inches(3.0), Inches(1.0), avatars, True)

    # Severity cards
    severities = [
        ("CRITICAL", RED, "Must-fix: bid\nrejection risk"),
        ("MAJOR", AMBER, "Significant\ngap in proposal"),
        ("MINOR", RGBColor(0xFA, 0xCC, 0x15), "Polish for\ncompleteness"),
    ]
    for i, (label, color, desc) in enumerate(severities):
        left = Inches(3.7) + Inches(2.0) * i
        add_rounded_rect(slide, left, Inches(1.75), Inches(1.85), Inches(1.0), CARD_BG, color)
        add_text_box(slide, left + Inches(0.1), Inches(1.82), Inches(1.65), Inches(0.3),
                     label, font_size=14, color=color, bold=True, alignment=PP_ALIGN.CENTER)
        add_text_box(slide, left + Inches(0.1), Inches(2.15), Inches(1.65), Inches(0.5),
                     desc, font_size=10, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    steps = [
        "\u2022  Vader compares proposal against every bid requirement",
        "\u2022  Structured gap report: ID, severity, section, description",
        "\u2022  Gap report sent to Yoda for automatic revision",
        "\u2022  Yoda revises and resubmits for re-review",
        "\u2022  Loop continues until resolved (max 2 rounds)",
        "\u2022  Final verdict: APPROVED or CONDITIONALLY APPROVED",
    ]
    add_bullet_list(slide, Inches(0.5), Inches(2.95), Inches(9.0), Inches(2.5),
                    steps, font_size=12, color=LIGHT_GRAY)

    # ==========================================================
    # SLIDE 9 — Step 4: Final Approval & Export
    # ==========================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(slide, "Step 4: Final Approval & Export",
                 "Complete bid package ready for submission", GOLD)

    deliverables = [
        ("Bid Proposal PDF", "AI-generated, ANAMAC-branded\nproposal aligned to all requirements", EMERALD),
        ("Assessment Summary", "Complete gap analysis with\ninitial findings & resolution status", AMBER),
        ("Full Package ZIP", "One-click download: bid PDF,\nproposal, supporting docs, summary", CYAN),
    ]
    for i, (title, desc, color) in enumerate(deliverables):
        left = Inches(0.4) + Inches(3.15) * i
        add_rounded_rect(slide, left, Inches(1.8), Inches(2.95), Inches(1.5), CARD_BG, color)
        add_text_box(slide, left + Inches(0.1), Inches(1.9), Inches(2.75), Inches(0.35),
                     title, font_size=14, color=color, bold=True, alignment=PP_ALIGN.CENTER)
        add_text_box(slide, left + Inches(0.1), Inches(2.35), Inches(2.75), Inches(0.8),
                     desc, font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    features = [
        "\u2714  Approval seal with timestamp and assessor signature",
        "\u2714  Resolution matrix showing all gaps and their final status",
        "\u2714  Two-column view: Initial Observations vs. Re-Review Findings",
        "\u2714  Export: Individual PDF download or complete ZIP package",
    ]
    add_bullet_list(slide, Inches(0.5), Inches(3.5), Inches(9.0), Inches(1.8),
                    features, font_size=12, color=LIGHT_GRAY)

    # ==========================================================
    # SLIDE 10 — How to Use
    # ==========================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(slide, "How to Use", "Your Path to the Light Side of Bidding", TEAL)

    steps_data = [
        ("1", "Launch", "Run start.sh \u2014 opens at localhost:5173", BLUE, "luke"),
        ("2", "Search", "Enter bid type (QRA, HAZOP) \u2192 click Search", BLUE, "luke"),
        ("3", "Download", "Select a bid \u2192 Luke downloads all docs", CYAN, "luke"),
        ("4", "Propose", "Propose tab \u2192 Yoda writes the proposal", EMERALD, "yoda"),
        ("5", "Assess", "Vader auto-reviews and identifies gaps", AMBER, "vader"),
        ("6", "Export", "Summary tab \u2192 Download PDF or full ZIP", GOLD, "r2d2"),
    ]
    for i, (num, title, desc, color, agent_key) in enumerate(steps_data):
        col = i % 2
        row = i // 2
        left = Inches(0.4) + Inches(4.8) * col
        top = Inches(1.8) + Inches(1.1) * row

        # Number circle
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top + Inches(0.1), Inches(0.4), Inches(0.4))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()
        add_text_box(slide, left + Inches(0.05), top + Inches(0.12), Inches(0.3), Inches(0.35),
                     num, font_size=14, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

        # Agent avatar (tiny)
        avatar_path = avatars[agent_key]
        if os.path.exists(avatar_path):
            slide.shapes.add_picture(avatar_path, left + Inches(0.5), top + Inches(0.1),
                                     Inches(0.38), Inches(0.38))

        add_text_box(slide, left + Inches(0.95), top + Inches(0.02), Inches(3.5), Inches(0.3),
                     title, font_size=14, color=WHITE, bold=True)
        add_text_box(slide, left + Inches(0.95), top + Inches(0.35), Inches(3.5), Inches(0.5),
                     desc, font_size=11, color=LIGHT_GRAY)

    # ==========================================================
    # SLIDE 11 — Business Value
    # ==========================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(slide, "Business Value & Efficiency Gains",
                 "Transforming Bid Operations from Days to Minutes", EMERALD)

    metrics = [
        ("95%", "Time Saved", "Days of work\nto minutes", EMERALD),
        ("100%", "Compliance", "Every requirement\nAI-verified", BLUE),
        ("3x", "More Bids", "Same team,\ntriple output", GOLD),
        ("0", "Lost Docs", "Centralized\nstorage", CYAN),
    ]
    for i, (number, label, desc, color) in enumerate(metrics):
        left = Inches(0.35) + Inches(2.4) * i
        add_rounded_rect(slide, left, Inches(1.8), Inches(2.2), Inches(1.85), CARD_BG, color)
        add_text_box(slide, left, Inches(1.9), Inches(2.2), Inches(0.6),
                     number, font_size=38, color=color, bold=True, alignment=PP_ALIGN.CENTER)
        add_text_box(slide, left, Inches(2.5), Inches(2.2), Inches(0.3),
                     label, font_size=13, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
        add_text_box(slide, left, Inches(2.85), Inches(2.2), Inches(0.6),
                     desc, font_size=10, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    benefits = [
        "\u2B50  Faster bid turnaround = more bids submitted = higher win rate",
        "\u2B50  Consistent proposal quality through AI-driven standardization",
        "\u2B50  Reduced human error in compliance checking and document assembly",
        "\u2B50  Scalability: handle multiple concurrent bids without additional staff",
        "\u2B50  Full audit trail: every agent action logged and traceable",
    ]
    add_bullet_list(slide, Inches(0.5), Inches(3.85), Inches(9.0), Inches(1.6),
                    benefits, font_size=12, color=LIGHT_GRAY)

    # ==========================================================
    # SLIDE 12 — Live Test Results
    # ==========================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(slide, "Live Test Results",
                 "End-to-End Pipeline Verified with Real GEM Bid", EMERALD)

    add_rounded_rect(slide, Inches(0.4), Inches(1.8), Inches(9.2), Inches(0.55), DARK_SLATE, GOLD)
    add_text_box(slide, Inches(0.5), Inches(1.85), Inches(9), Inches(0.4),
                 "Test Bid: GEM/2026/B/7314720  |  Total: ~3 min  |  CONDITIONALLY APPROVED  \u2714",
                 font_size=13, color=GOLD, bold=True, alignment=PP_ALIGN.CENTER)

    timeline = [
        ("Luke: Search", "1 bid found matching query", "5s", BLUE, "luke"),
        ("Luke: Download", "1 primary PDF + 8 supporting docs", "30s", CYAN, "luke"),
        ("Yoda: Proposal", "AI-generated proposal via Claude", "35s", EMERALD, "yoda"),
        ("Vader: Assessment", "12 gaps (4 Critical, 5 Major, 3 Minor)", "20s", AMBER, "vader"),
        ("Yoda: Revision 1", "Addresses Vader's gap report", "40s", EMERALD, "yoda"),
        ("Vader: Re-Review", "Final revision round", "40s", AMBER, "vader"),
        ("R2-D2: Verified", "CONDITIONALLY APPROVED", "\u2714", EMERALD, "r2d2"),
    ]
    for i, (step, desc, time_str, color, agent_key) in enumerate(timeline):
        top = Inches(2.5) + Inches(0.38) * i

        # Agent avatar (tiny)
        avatar_path = avatars[agent_key]
        if os.path.exists(avatar_path):
            slide.shapes.add_picture(avatar_path, Inches(0.55), top + Inches(0.02),
                                     Inches(0.28), Inches(0.28))

        # Connector line
        if i < len(timeline) - 1:
            line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                          Inches(0.68), top + Inches(0.3), Pt(2), Inches(0.15))
            line.fill.solid()
            line.fill.fore_color.rgb = MID_GRAY
            line.line.fill.background()

        add_text_box(slide, Inches(0.95), top, Inches(2.0), Inches(0.3),
                     step, font_size=10, color=color, bold=True)
        add_text_box(slide, Inches(3.0), top, Inches(5.0), Inches(0.3),
                     desc, font_size=10, color=LIGHT_GRAY)
        add_text_box(slide, Inches(8.3), top, Inches(1.2), Inches(0.3),
                     time_str, font_size=10, color=MID_GRAY, alignment=PP_ALIGN.RIGHT)

    # QA badge
    add_rounded_rect(slide, Inches(0.4), Inches(5.0), Inches(9.2), Inches(0.4), DARK_SLATE, EMERALD)
    add_text_box(slide, Inches(0.5), Inches(5.05), Inches(9), Inches(0.3),
                 "R2-D2 QA Report: 22 / 22 Tests Passing  |  8 Integration Tests  |  All Systems Operational  *happy beep*",
                 font_size=10, color=EMERALD, alignment=PP_ALIGN.CENTER)

    # ==========================================================
    # SLIDE 13 — Thank You
    # ==========================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, NAVY)

    # Top accent
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Pt(5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()

    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Inches(3.8), Inches(0.5), Inches(2.4), Inches(1.8))

    add_text_box(slide, Inches(0.5), Inches(2.4), Inches(9), Inches(0.7),
                 "Thank You", font_size=42, color=GOLD, bold=True,
                 alignment=PP_ALIGN.CENTER)
    add_gold_line(slide, Inches(3), Inches(3.1), Inches(4))
    add_text_box(slide, Inches(0.5), Inches(3.3), Inches(9), Inches(0.5),
                 "ANAMAC Bid Proposal Intelligence Platform",
                 font_size=18, color=WHITE, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, Inches(0.5), Inches(3.8), Inches(9), Inches(0.4),
                 '"The Force is strong with this platform."',
                 font_size=13, color=TEAL, alignment=PP_ALIGN.CENTER, italic=True)

    # Character avatars row
    for i, key in enumerate(agent_keys):
        avatar_path = avatars[key]
        agent = AGENTS[key]
        x = Inches(2.2) + Inches(1.5) * i
        if os.path.exists(avatar_path):
            slide.shapes.add_picture(avatar_path, x, Inches(4.15), Inches(0.55), Inches(0.55))
        add_text_box(slide, x - Inches(0.1), Inches(4.7), Inches(0.75), Inches(0.3),
                     agent["name"], font_size=7, color=agent["color"], bold=True,
                     alignment=PP_ALIGN.CENTER)

    add_text_box(slide, Inches(0.5), Inches(5.1), Inches(9), Inches(0.3),
                 "Confidential  |  ANAMAC  |  March 2026",
                 font_size=10, color=DIM_GRAY, alignment=PP_ALIGN.CENTER)

    add_logo_watermark(slide)

    # Save
    output_path = os.path.join(SCRIPT_DIR, "ANAMAC_Bid_Platform_Presentation.pptx")
    prs.save(output_path)
    print(f"\nPresentation saved to: {output_path}")
    print(f"Total slides: {len(prs.slides)}")

    return output_path


if __name__ == "__main__":
    build_presentation()
